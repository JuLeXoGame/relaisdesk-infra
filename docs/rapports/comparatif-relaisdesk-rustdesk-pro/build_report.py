from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "Comparatif_RelaisDesk_RustDesk_Pro_2026-08-26.docx"

# Selected design system: standard_business_brief + memo_masthead.
PAGE_W_DXA = 12240
PAGE_H_DXA = 15840
CONTENT_W_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_START_END_DXA = 120

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B7A"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
INK = "1F2937"
GREEN = "1F5D42"
GOLD = "7A5A00"
RED = "9B1C1C"
BORDER = "C9D2DC"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != CONTENT_W_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_W_DXA}, got {sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, CELL_TOP_BOTTOM_DXA, CELL_START_END_DXA, CELL_TOP_BOTTOM_DXA, CELL_START_END_DXA)


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, fld_sep, text, fld_end])
    set_run_font(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color=BLUE, size="18", space="5") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def create_numbering(doc: Document) -> tuple[int, int, int]:
    numbering = doc.part.numbering_part.element

    existing_abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    existing_num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    bullet_abstract_id = max(existing_abstract_ids, default=0) + 1
    decimal_abstract_id = bullet_abstract_id + 1
    priorities_abstract_id = decimal_abstract_id + 1
    bullet_num_id = max(existing_num_ids, default=0) + 1
    findings_num_id = bullet_num_id + 1
    priorities_num_id = bullet_num_id + 2

    def make_abstract(fmt: str, text: str, abstract_id: int) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"A1B2C{abstract_id:03d}")
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Calibri")
        r_fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(r_fonts)
        lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def make_instance(abstract_id: int, num_id: int) -> None:
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id_el = OxmlElement("w:abstractNumId")
        abstract_id_el.set(qn("w:val"), str(abstract_id))
        num.append(abstract_id_el)
        lvl_override = OxmlElement("w:lvlOverride")
        lvl_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        lvl_override.append(start_override)
        num.append(lvl_override)
        numbering.append(num)

    make_abstract("bullet", "•", bullet_abstract_id)
    make_abstract("decimal", "%1.", decimal_abstract_id)
    make_abstract("decimal", "%1.", priorities_abstract_id)
    make_instance(bullet_abstract_id, bullet_num_id)
    make_instance(decimal_abstract_id, findings_num_id)
    make_instance(priorities_abstract_id, priorities_num_id)
    return bullet_num_id, findings_num_id, priorities_num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for custom_name, size, color, before, after, line in (
        ("Source Note", 9, MUTED, 4, 4, 1.0),
        ("Table Text", 8.2, INK, 0, 2, 1.0),
        ("Table Header", 8.4, NAVY, 0, 0, 1.0),
        ("Callout", 11, NAVY, 0, 0, 1.12),
    ):
        if custom_name not in styles:
            style = styles.add_style(custom_name, 1)
        else:
            style = styles[custom_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False


def set_running_header_footer(section) -> None:
    def populate_footer(footer) -> None:
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, "PAGE")

    # Keep only the page number in the footer. The report uses a continuous
    # flow after its cover to avoid Word PDF-export quirks at forced breaks.
    populate_footer(section.footer)


def add_kicker(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    set_run_font(run, size=9.5, color=BLUE, bold=True)


def add_title_block(doc: Document) -> None:
    add_kicker(doc, "Comparatif stratégique")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("RelaisDesk face aux offres payantes RustDesk")
    set_run_font(run, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Avantages réels, limites concurrentielles et positionnement recommandé")
    set_run_font(run, size=13.5, color=MUTED)

    metadata = [
        ("Objet", "Définir ce que RelaisDesk apporte réellement de plus"),
        ("Périmètre", "Projet local actuel vs. documentation officielle RustDesk consultée le 26 août 2026"),
        ("Statut", "Analyse de préproduction — décision de positionnement"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label} : ")
        set_run_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(12)
    add_bottom_border(rule)


def add_callout(doc: Document, lead: str, body: str, fill=CALLOUT, accent=BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W_DXA])
    set_repeat_table_header(table.rows[0])
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.style = doc.styles["Callout"]
    r = p.add_run(lead + " ")
    set_run_font(r, size=11, color=NAVY, bold=True)
    r = p.add_run(body)
    set_run_font(r, size=11, color=NAVY)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_para(doc: Document, text: str, bold_prefix: str | None = None, style=None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullet(doc: Document, text: str, bullet_id: int, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    apply_num(p, bullet_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_compact_bullet(doc: Document, text: str, bullet_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, bullet_id)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=9.5, color=INK)


def add_number(doc: Document, text: str, decimal_id: int, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    apply_num(p, decimal_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Source Note")
    p.add_run(text)


def add_matrix(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]], widths: Sequence[int]) -> None:
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Header"]
        p.paragraph_format.keep_with_next = True
        r = p.add_run(header)
        set_run_font(r, size=8.4, color=NAVY, bold=True)

    for row_idx, values in enumerate(rows):
        row = table.add_row()
        if row_idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, "FAFBFC")
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            r = p.add_run(value)
            set_run_font(r, size=8.2, color=INK, bold=(col_idx == 0))
    set_table_geometry(table, widths)


def add_page_heading(doc: Document, text: str, level: int = 1):
    """Add a major heading in the continuous report flow."""
    return doc.add_heading(text, level=level)


def add_source_entry(doc: Document, code: str, label: str, url: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    r = p.add_run(f"{code} — {label}. ")
    set_run_font(r, size=9, color=INK, bold=True)
    add_hyperlink(p, "Consulter la source officielle", url)


def audit_geometry(doc: Document) -> None:
    section = doc.sections[0]
    expected = {
        "page_width": PAGE_W_DXA,
        "page_height": PAGE_H_DXA,
        "top_margin": 1440,
        "right_margin": 1440,
        "bottom_margin": 1440,
        "left_margin": 1440,
    }
    for attr, value in expected.items():
        actual = int(getattr(section, attr).twips)
        if actual != value:
            raise AssertionError(f"{attr}: expected {value}, got {actual}")
    for table in doc.tables:
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        if tbl_w is None or tbl_w.get(qn("w:w")) != str(CONTENT_W_DXA):
            raise AssertionError("A body table does not use fixed 9360 DXA width")


def build() -> Path:
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    for section in doc.sections:
        configure_section(section)
        set_running_header_footer(section)
    configure_styles(doc)
    bullet_id, findings_num_id, priorities_num_id = create_numbering(doc)

    core = doc.core_properties
    core.title = "RelaisDesk face aux offres payantes RustDesk"
    core.subject = "Comparatif stratégique et positionnement produit"
    core.author = "RelaisDesk — Informatique à domicile 03"
    core.keywords = "RelaisDesk, RustDesk Pro, assistance à distance, comparatif"
    core.comments = "Étude fondée sur le projet local et les sources officielles RustDesk consultées le 26 août 2026."

    add_title_block(doc)
    add_callout(
        doc,
        "Verdict.",
        "RelaisDesk n’est pas plus complet que RustDesk Pro en général. Il est mieux verticalisé pour vendre et exploiter des interventions d’assistance à distance : client occasionnel, capacité technicien simultanée, commerce, renouvellement et historique métier.",
    )

    doc.add_heading("Réponse en une phrase", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("RelaisDesk transforme le socle libre RustDesk en service français commercialisable : ")
    set_run_font(r, size=12, color=NAVY, bold=True)
    r = p.add_run("codes Viewer temporaires, accès isolé par client, facturation, espace client et dossiers d’intervention — sans facturer chaque PC client installé.")
    set_run_font(r, size=12, color=NAVY)

    doc.add_heading("Le point de distinction", level=1)
    add_para(doc, "RustDesk Pro est un produit d’administration centralisée d’utilisateurs et d’appareils. RelaisDesk est un produit d’exploitation commerciale d’un service de dépannage à distance. Les deux se recoupent sur le transport RustDesk et le support occasionnel, mais ils n’optimisent pas le même métier.")
    add_source_note(doc, "Base de comparaison : projet local RelaisDesk [L1–L5] ; tarification et documentation officielles RustDesk [S1–S12].")

    doc.add_page_break()
    add_page_heading(doc, "1. Conclusions décisives")
    findings = [
        ("Le différenciateur principal est commercial.", " RelaisDesk relie achat, paiement, facture, licence, renouvellement, rappel et espace client. Le bureau à distance reste fourni par le socle RustDesk forké."),
        ("Le modèle de licence correspond aux prestataires.", " Il limite les appareils techniciens actifs simultanément, pas les installations inactives. Les Viewers temporaires sont annoncés comme illimités et chaque code expire après douze heures."),
        ("Le parcours occasionnel est intégré de bout en bout.", " Un code Viewer configure l’accès, lie le premier appareil, isole le locataire et ouvre automatiquement une fiche d’intervention."),
        ("RustDesk Customized V2 chevauche déjà cet usage.", " Les appareils excédentaires ne sont pas désactivés et les canaux simultanés peuvent être limités. L’argument “nous seuls savons faire du support ponctuel” serait donc faux. [S1, S5, S12]"),
        ("La sécurité RelaisDesk est spécialisée, pas supérieure en tout.", " Jetons courts, preuve de possession, anti-rejeu et séparation technicien/Viewer répondent au risque de partage de licence ; Pro conserve un avantage sur la 2FA, l’OIDC, LDAP, les politiques et l’audit. [S2, S7–S10]"),
        ("RelaisDesk est encore en préproduction.", " Les sources des forks doivent être publiées, les exécutables signés et les recettes P2P/relais/révocation/quota terminées avant toute promesse de maturité. [L5]"),
    ]
    for lead, tail in findings:
        add_number(doc, lead + tail, findings_num_id, bold_prefix=lead)

    add_callout(
        doc,
        "Positionnement recommandé.",
        "« Plateforme française de commercialisation et d’exploitation de l’assistance à distance, basée sur RustDesk libre. »",
        fill=PALE_BLUE,
        accent=DARK_BLUE,
    )

    add_page_heading(doc, "2. Comparaison — métier et expérience client")
    business_rows = [
        ("Positionnement", "Service d’assistance à vendre à des techniciens et petites équipes.", "Administration auto-hébergée d’utilisateurs et d’appareils.", "Différent"),
        ("Commerce", "Souscription, Stripe, virement, factures, B2B/B2C et CGV.", "Licence vendue à l’exploitant ; pas de pile de revente aval présentée.", "RelaisDesk"),
        ("Portail", "Espace client : licences, commandes, factures, renouvellements et interventions.", "Console Web : utilisateurs, appareils, permissions, politiques, journaux et relais.", "Différent"),
        ("Unité de licence", "Postes techniciens actifs simultanément ; installations inactives non décomptées.", "Plans classiques par utilisateurs/appareils ; V2 par canaux simultanés avec quotas de gestion.", "Mixte"),
        ("Support ponctuel", "Code Viewer 12 h, sans compte, lié au premier appareil et à une intervention.", "Quick Support et Custom2 ; appareil non enregistré possible avec restrictions.", "RelaisDesk"),
        ("Historique", "Dossier métier : client, objet, durée, statut, compte rendu et export CSV.", "Audit technique : connexions, fichiers, alarmes et opérations console.", "Différent"),
        ("Renouvellement", "Renouvellement mensuel guidé ; rappels J-14/J-7/J-3/J-1 et après expiration.", "Paiement annuel, pas d’auto-renouvellement ; rappel 14 jours avant expiration.", "RelaisDesk"),
        ("Marque", "Expérience et contrats RelaisDesk destinés au marché français.", "Clients personnalisables : nom, logo, icône, configuration et signature.", "Mixte"),
    ]
    add_matrix(doc, ["Domaine", "RelaisDesk", "RustDesk Pro", "Lecture"], business_rows, [1550, 3100, 3100, 1610])
    add_source_note(doc, "Sources : code et documentation locale [L1–L5] ; tarification, console, configuration et modèle casual support RustDesk [S1, S3–S6, S12].")

    doc.add_heading("Interprétation", level=2)
    add_para(doc, "Le portail RelaisDesk et la console RustDesk Pro ne sont pas des doublons. Le premier s’adresse au client commercial qui achète le service ; la seconde s’adresse à l’administrateur technique qui gère des comptes, des appareils et des politiques.")

    add_page_heading(doc, "3. Sécurité, parc et exploitation")
    technical_rows = [
        ("Autorisation réseau", "Jeton Ed25519 de 5 min, preuve d’appareil, nonce anti-rejeu, rôle et locataire vérifiés dans hbbs/hbbr.", "Comptes et règles centralisées ; contrôles Pro appliqués à l’écosystème géré.", "Spécialisé"),
        ("Isolation", "Technicien → Viewer du même locataire uniquement ; code Viewer lié au premier appareil.", "ACL par utilisateurs, groupes d’utilisateurs et groupes d’appareils.", "Mixte"),
        ("Identité", "Licence technicien ; lien magique pour le client commercial.", "2FA e-mail/TOTP, OIDC, LDAP et codes de secours.", "RustDesk"),
        ("Parc", "Pas d’équivalent complet aux inventaires et groupes Pro.", "Appareils, groupes, affectations et carnets d’adresses.", "RustDesk"),
        ("Politiques", "Configuration RelaisDesk imposée par les launchers ; contrôle de licence dans le fork.", "Stratégies centralisées et rôles de contrôle pendant la session.", "RustDesk"),
        ("Audit", "Historique métier volontairement limité ; aucun contenu de session copié.", "Audits connexion, fichier, alarme et console avec filtres/API.", "RustDesk"),
        ("Déploiement", "Launchers Windows ; Linux prévu ; chaîne de publication encore à finaliser.", "Générateur multi-plateforme, MSI, scripts, RMM, Intune et GPO.", "RustDesk"),
        ("Relais / Web", "Un hbbs/hbbr direct ; WebSocket fermé tant qu’inutile.", "Relais distribués avec sélection du plus proche ; client Web selon licence.", "RustDesk"),
        ("Maintenance", "Contrôle de la feuille de route, avec charge AGPL, CI, correctifs et support.", "Maintenance éditeur et support dédié inclus dans les plans payants.", "RustDesk"),
    ]
    add_matrix(doc, ["Domaine", "RelaisDesk", "RustDesk Pro", "Lecture"], technical_rows, [1550, 3100, 3100, 1610])
    add_source_note(doc, "Sources : architecture et sécurité locales [L2–L5] ; Server Pro, console, 2FA, ACL, rôles et stratégies [S2, S3, S7–S10].")

    add_page_heading(doc, "4. Les cinq avantages défendables de RelaisDesk")
    advantages = [
        ("Chaîne commerciale intégrée", "Le projet traite la commande, le paiement, la facture, la licence et le renouvellement dans le même modèle de données. Le prix est recalculé côté serveur et les événements Stripe sont signés et idempotents."),
        ("Licence alignée sur le travail réel", "Un petit prestataire peut installer le technicien sur plusieurs ordinateurs et ne consommer que les appareils actifs simultanément. C’est plus facile à vendre qu’un inventaire d’appareils gérés lorsque les clients sont occasionnels."),
        ("Parcours Viewer temporaire", "Le code de douze heures ne sert pas seulement de mot de passe : il établit le locataire, lie l’appareil et déclenche la fiche professionnelle d’intervention."),
        ("Isolation anti-partage dans le réseau", "Le fork refuse les connexions sans autorisation signée, impose les rôles complémentaires et le même locataire, puis propage expiration et révocation grâce aux jetons courts et au canal de santé."),
        ("Restitution métier au client", "Le client commercial retrouve factures, renouvellements et compte rendu de ses interventions dans un espace distinct de l’outil technicien."),
    ]
    for lead, body in advantages:
        doc.add_heading(lead, level=2)
        add_para(doc, body)

    add_callout(
        doc,
        "Nuance essentielle.",
        "RustDesk sait déjà faire du Quick Support et propose Customized V2. L’avantage RelaisDesk vient de l’orchestration commerciale et de la continuité code Viewer → autorisation → intervention → espace client.",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_page_heading(doc, "5. Tarification et modèle d’usage")
    doc.add_heading("RelaisDesk — prix encodés dans le projet", level=2)
    rd_prices = [
        ("Starter", "10 € / 30 jours", "1 poste technicien actif", "Viewers 12 h annoncés illimités"),
        ("Pro", "20 € / 30 jours", "jusqu’à 10 actifs", "Viewers 12 h annoncés illimités"),
        ("Ultra", "par tranches", "au-delà de 10", "Barème progressif côté serveur"),
    ]
    add_matrix(doc, ["Offre", "Prix", "Capacité", "Viewer"], rd_prices, [1500, 1900, 2600, 3360])
    add_source_note(doc, "Source locale : `database/orders.go` et `api/handlers/public_orders.go` [L4].")

    doc.add_heading("RustDesk Server Pro — prix affichés le 26 août 2026", level=2)
    pro_prices = [
        ("Individual", "11,88 $/mois, annuel", "1", "20", "illimitées"),
        ("Basic", "23,88 $/mois, annuel", "10", "100", "illimitées"),
        ("Customized", "23,88 $ + suppléments", "+ 1,20 $", "+ 0,12 $", "illimitées"),
        ("Customized V2", "23,88 $ + suppléments", "+ 1,20 $", "+ 0,12 $", "+ 24 $ / canal"),
    ]
    add_matrix(doc, ["Offre", "Prix affiché", "Utilisateur", "Appareil", "Concurrence"], pro_prices, [1450, 2500, 1650, 1650, 2110])
    add_source_note(doc, "Source officielle : page de tarification RustDesk consultée le 26 août 2026 [S1]. Prix en USD, facturation annuelle.")

    doc.add_heading("Ce que cette comparaison permet — et ne permet pas", level=2)
    add_bullet(doc, "RelaisDesk est commercialement simple pour des techniciens utilisant beaucoup de PC possibles mais peu de postes simultanément.", bullet_id)
    add_bullet(doc, "Customized V2 rend RustDesk compétitif sur le même cas d’usage ; les appareils excédentaires ne sont pas désactivés mais restent limités pour les affectations, groupes et carnets d’adresses. [S12]", bullet_id)
    add_bullet(doc, "Le prix RelaisDesk ne peut pas être déclaré universellement inférieur sans intégrer serveur, trafic de relais, sauvegardes, signature de code, support, maintenance et conformité.", bullet_id)
    add_bullet(doc, "Le terme exact est “renouvellement guidé avec relances automatiques” : aucun prélèvement ni renouvellement automatique n’est implémenté aujourd’hui.", bullet_id)

    add_page_heading(doc, "6. Là où RustDesk Pro reste nettement supérieur")
    pro_strengths = [
        "Gestion centralisée des utilisateurs, groupes, appareils, groupes d’appareils et carnets d’adresses.",
        "2FA e-mail/TOTP, OIDC, LDAP, codes de secours et rôles administratifs délégués.",
        "Contrôles d’accès croisés, rôles définissant les actions possibles pendant une session et stratégies de sécurité synchronisées.",
        "Audit technique des connexions, transferts de fichiers, alarmes et opérations de console.",
        "Relais multiples avec sélection automatique du plus proche.",
        "Générateur de client personnalisé et déploiement documenté pour Windows, macOS, Linux, Android, MSI, RMM, Intune et GPO.",
        "Client Web selon les conditions de licence, support dédié inclus et maturité commerciale supérieure.",
    ]
    for item in pro_strengths:
        add_bullet(doc, item, bullet_id)

    add_callout(
        doc,
        "Choix produit.",
        "Pour une DSI interne qui administre un parc permanent et exige SSO, politiques et audit, RustDesk Pro est actuellement plus complet. Pour un prestataire qui facture des interventions ponctuelles à de nombreux clients, RelaisDesk est mieux ajusté au métier.",
    )

    doc.add_heading("Principal manque de sécurité de compte", level=2)
    add_para(doc, "RelaisDesk possède des protections réseau spécialisées, mais ne propose pas encore l’équivalent de la 2FA centralisée RustDesk Pro pour les comptes administrateur et technicien. Avant d’affirmer une sécurité supérieure, la priorité raisonnable est d’ajouter une 2FA au moins à ces comptes sensibles.")

    add_page_heading(doc, "7. Positionnement commercial recommandé")
    doc.add_heading("Promesse centrale", level=2)
    add_callout(
        doc,
        "RelaisDesk permet aux professionnels de vendre, sécuriser et tracer leurs interventions à distance,",
        "sans facturer chaque PC client installé.",
        fill=PALE_BLUE,
        accent=DARK_BLUE,
    )

    doc.add_heading("Formulations solides", level=2)
    solid = [
        "Une solution d’assistance à distance pensée pour les techniciens indépendants et les équipes de support.",
        "Installez l’application technicien sur plusieurs postes ; votre offre limite uniquement les postes actifs simultanément.",
        "Donnez à votre client un code Viewer temporaire de douze heures, sans création de compte.",
        "Retrouvez licences, factures, renouvellements et historiques d’intervention dans un espace client dédié.",
        "Les connexions sont isolées par client et autorisées par des jetons courts liés à l’appareil.",
        "Basé sur RustDesk libre, avec flux direct P2P ou relais lorsque nécessaire.",
    ]
    for item in solid:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("Formulations à éviter", level=2)
    avoid = [
        "“Plus sécurisé que RustDesk Pro” — non démontrable globalement.",
        "“Plus complet que RustDesk Pro” — faux pour le parc, l’identité et les politiques.",
        "“Appareils illimités” sans préciser Viewers temporaires ou installations inactives.",
        "“Renouvellement automatique” — la version actuelle exige une action et un paiement exprès.",
        "“Aucun équivalent chez RustDesk” pour le support ponctuel ou la concurrence — Customized V2 existe.",
        "“Prêt pour la production” avant publication des sources, signature des exécutables et recette complète.",
    ]
    for item in avoid:
        add_bullet(doc, item, bullet_id)

    add_page_heading(doc, "8. Segments cibles et priorités")
    segments = [
        ("Cœur de cible", "Techniciens indépendants, artisans informatiques, TPE de dépannage et petites équipes externes servant beaucoup de clients occasionnels."),
        ("Cible secondaire", "MSP légers, associations et organismes ayant besoin de commercialiser ou tracer des créneaux d’assistance ponctuelle."),
        ("Mauvais ajustement actuel", "DSI exigeant OIDC/LDAP, parc permanent, politiques de sécurité centralisées, client Web, relais géographiques et audit exhaustif."),
    ]
    add_matrix(doc, ["Segment", "Lecture"], segments, [2200, 7160])

    doc.add_heading("Priorités issues du comparatif", level=2)
    priorities = [
        ("Terminer la mise en production.", " Publier les forks et sources correspondantes, exécuter la CI, produire Linux, signer Windows et réaliser la recette réseau complète."),
        ("Ajouter la 2FA aux comptes sensibles.", " Commencer par l’administration et les techniciens."),
        ("Renforcer l’historique métier.", " Validation du compte rendu, PDF d’intervention et filtres par période/client, sans chercher à recopier tout l’audit Pro."),
        ("Définir publiquement l’unité de licence.", " Expliquer “poste technicien actif”, la libération d’un créneau et la différence entre Viewer temporaire et appareil géré."),
        ("Éviter l’inventaire de parc par défaut.", " Ne construire cet ensemble coûteux que si le marché le réclame ; il placerait RelaisDesk face à la force principale de RustDesk Pro."),
    ]
    for lead, tail in priorities:
        add_number(doc, lead + tail, priorities_num_id, bold_prefix=lead)

    # Let the methodological appendix follow the priorities on the same page.
    # This avoids stranding a memo masthead as the final Word page is repaginated.
    doc.add_heading("9. Méthode, confiance et limites", level=1)
    add_para(doc, "L’étude combine une lecture du code et de la documentation locale RelaisDesk avec deux vagues de recherche sur les sources officielles RustDesk : d’abord les offres et fonctions générales, puis les points susceptibles d’invalider une différenciation (Customized V2, Quick Support, `register-device=N`, définitions des connexions simultanées, 2FA, rôles et stratégies).")
    add_para(doc, "Confiance élevée sur les fonctions et tarifs RustDesk cités, consultés le jour de l’étude. Confiance élevée sur les fonctions RelaisDesk vérifiées dans le projet local. Confiance moyenne sur l’absence de fonctions commerciales chez RustDesk : elle signifie qu’elles ne sont pas présentées dans les pages officielles étudiées, sans exclure une intégration tierce ou privée.")
    add_para(doc, "Le comparatif porte sur le produit et le positionnement. Il ne constitue ni une consultation juridique, ni une certification de sécurité, ni une validation de production.")

    doc.add_heading("Sources officielles RustDesk", level=2)
    sources = [
        ("S1", "Tarification et FAQ commerciale", "https://rustdesk.com/pricing/?lang=en"),
        ("S2", "Vue d’ensemble de RustDesk Server Pro", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/"),
        ("S3", "Console Web", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/console/"),
        ("S4", "Configuration et génération de clients", "https://rustdesk.com/docs/en/self-host/client-configuration/"),
        ("S5", "Paramètres avancés, dont register-device=N", "https://rustdesk.com/docs/en/self-host/client-configuration/advanced-settings/"),
        ("S6", "Déploiement des clients", "https://rustdesk.com/docs/en/self-host/client-deployment/"),
        ("S7", "Authentification 2FA", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/2fa/"),
        ("S8", "Contrôle d’accès", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/permissions/"),
        ("S9", "Rôles de contrôle", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/control-role/"),
        ("S10", "Stratégies centralisées", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/strategy/"),
        ("S11", "LDAP", "https://rustdesk.com/docs/en/self-host/rustdesk-server-pro/ldap/"),
        ("S12", "Définitions de licence et casual support", "https://github.com/rustdesk/rustdesk/wiki/FAQ#what-are-managed-devices--login-users--concurrent-connections-in-pro"),
        ("S13", "Réponse officielle sur Customized V2", "https://github.com/rustdesk/rustdesk-server-pro/discussions/182"),
    ]
    for code, label, url in sources:
        add_source_entry(doc, code, label, url)

    doc.add_heading("Sources locales RelaisDesk", level=2)
    local_sources = [
        "L1 — README.md",
        "L2 — docs/ARCHITECTURE.md et docs/SECURITY.md",
        "L3 — docs/FORK_AUTHORIZATION.md",
        "L4 — database/orders.go, database/viewer.go, database/customer_commercial.go et handlers API",
        "L5 — docs/FORK_IMPLEMENTATION_STATUS.md",
    ]
    for item in local_sources:
        add_compact_bullet(doc, item, bullet_id)

    audit_geometry(doc)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
